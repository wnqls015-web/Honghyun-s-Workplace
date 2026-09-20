"""
회의 전사(transcript)를 받아 Claude API로 회의록 항목을 구조화 추출하고,
사내 표준 회의록 양식(template.xlsx)을 그대로 채운 .xlsx 리포트(RP)를 생성한 뒤
Supabase Storage에 업로드 + rp_reports 테이블에 메타데이터를 저장하는 스크립트.

사용 예:
    python generate_rp.py --input transcript.txt --title "주간업무보고 회의"

환경변수 (.env):
    ANTHROPIC_API_KEY
    SUPABASE_URL
    SUPABASE_SERVICE_KEY
"""

import os
import re
import json
import argparse
import smtplib
from copy import copy
from datetime import date
from email.message import EmailMessage

import anthropic
import openpyxl
from supabase import create_client

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "template.xlsx")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")
STORAGE_BUCKET = "rp-reports"

# 템플릿의 고정 좌표 (automation/rp/template.xlsx 기준)
# 1~10행(제목/회의개요/표 헤더)은 항상 고정이며, 11행부터를 항목 수에 맞춰 새로 그린다.
REBUILD_FROM_ROW = 11
AGENDA_BLOCK_ROWS = 4        # 블록 1개(순번/안건/논의내용)가 차지하는 행 수
AGENDA_BLOCK_HEIGHT = 199.5
SPACER_ROW_TEMPLATE = 15
INSTR_HEADER_ROW_TEMPLATE = 16
INSTR_TABLE_HEADER_ROW_TEMPLATE = 17
INSTR_ROW_TEMPLATE = 18
INSTR_ROW_HEIGHT = 30


def safe_filename(name: str) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣_-]+", "_", name) or "RP"


def call_claude_for_structure(transcript: str) -> dict:
    client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])

    message = client.messages.create(
        model="claude-sonnet-5",
        max_tokens=4000,
        system=(
            "너는 회의 전사록을 읽고 사내 표준 회의록 양식에 맞춰 항목을 구조화하는 비서야. "
            "원문에 없는 내용을 지어내지 말고, 불확실한 항목은 그대로 '확인 필요'라고 표기해. "
            "반드시 아래 JSON 스키마 형식으로만, 다른 설명 없이 응답해:\n"
            "{\n"
            '  "meeting_name": "회의명",\n'
            '  "datetime": "일시",\n'
            '  "location": "장소",\n'
            '  "author": "작성자",\n'
            '  "attendees": "참석자 (명단/역할)",\n'
            '  "agenda_items": [{"agenda": "안건", "discussion": "논의 내용 (핵심만 불릿 형태로)"}],\n'
            '  "instructions": [{"instruction": "지시사항", "owner": "담당자", "note": "비고(기한 등)"}]\n'
            "}"
        ),
        messages=[
            {
                "role": "user",
                "content": f"다음 회의 전사록을 정리해줘:\n\n{transcript}",
            }
        ],
    )
    raw = message.content[0].text.strip()
    raw = re.sub(r"^```(json)?|```$", "", raw, flags=re.MULTILINE).strip()
    return json.loads(raw)


def build_raw_data(transcript: str, meeting_date: str) -> dict:
    """ANTHROPIC_API_KEY가 없을 때: AI 정리 없이 전사록 원문을 그대로 담는다."""
    return {
        "meeting_name": f"{meeting_date} 회의",
        "datetime": meeting_date,
        "location": "",
        "author": "",
        "attendees": "",
        "agenda_items": [{"agenda": "전사록 원문", "discussion": transcript}],
        "instructions": [],
    }


def capture_block_style(ws, rows, cols):
    styles = {}
    for i, r in enumerate(rows):
        for c in cols:
            cell = ws.cell(row=r, column=c)
            styles[(i, c)] = {
                "font": copy(cell.font),
                "border": copy(cell.border),
                "fill": copy(cell.fill),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
            }
    return styles


def apply_block_style(ws, start_row, styles):
    for (offset, c), st in styles.items():
        cell = ws.cell(row=start_row + offset, column=c)
        cell.font = st["font"]
        cell.border = st["border"]
        cell.fill = st["fill"]
        cell.alignment = st["alignment"]
        cell.number_format = st["number_format"]
        cell.protection = st["protection"]


def capture_row_values(ws, row, cols):
    return {c: ws.cell(row=row, column=c).value for c in cols}


def clear_rows_from(ws, start_row):
    """start_row부터 시트 끝까지의 병합/행높이/값을 모두 제거해 새로 그릴 수 있게 비운다."""
    for r in list(ws.merged_cells.ranges):
        if r.min_row >= start_row:
            ws.unmerge_cells(str(r))
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)
    for r in [r for r in ws.row_dimensions.keys() if r >= start_row]:
        del ws.row_dimensions[r]


def fill_template(data: dict, output_path: str):
    # 스타일 캡처용으로 원본 템플릿을 별도로 한 번 더 읽는다(수정용 wb와 분리).
    tmpl_wb = openpyxl.load_workbook(TEMPLATE_PATH)
    tmpl_ws = tmpl_wb["회의록"]

    cols = range(2, 10)  # B~I
    agenda_style = capture_block_style(tmpl_ws, rows=[11, 12, 13, 14], cols=cols)
    spacer_style = capture_block_style(tmpl_ws, rows=[SPACER_ROW_TEMPLATE], cols=cols)
    instr_header_style = capture_block_style(tmpl_ws, rows=[INSTR_HEADER_ROW_TEMPLATE], cols=cols)
    instr_header_values = capture_row_values(tmpl_ws, INSTR_HEADER_ROW_TEMPLATE, cols)
    instr_table_header_style = capture_block_style(tmpl_ws, rows=[INSTR_TABLE_HEADER_ROW_TEMPLATE], cols=cols)
    instr_table_header_values = capture_row_values(tmpl_ws, INSTR_TABLE_HEADER_ROW_TEMPLATE, cols)
    instr_row_style = capture_block_style(tmpl_ws, rows=[INSTR_ROW_TEMPLATE], cols=cols)

    agenda_items = data.get("agenda_items") or [{"agenda": "", "discussion": ""}]
    instructions = data.get("instructions") or [{"instruction": "해당 없음", "owner": "", "note": ""}]

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb["회의록"]

    # 1~10행(제목/회의개요/표 헤더)은 그대로 두고, 11행 이후는 항목 수에 맞춰 새로 그린다.
    clear_rows_from(ws, REBUILD_FROM_ROW)

    cursor = REBUILD_FROM_ROW

    # 회의 내용: 안건 블록 (순번/안건/논의내용, 4행 병합)
    for k, item in enumerate(agenda_items):
        r = cursor
        apply_block_style(ws, r, agenda_style)
        ws.merge_cells(start_row=r, start_column=2, end_row=r + 3, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r + 3, end_column=4)
        ws.merge_cells(start_row=r, start_column=5, end_row=r + 3, end_column=9)
        for rr in range(r, r + AGENDA_BLOCK_ROWS):
            ws.row_dimensions[rr].height = AGENDA_BLOCK_HEIGHT
        ws.cell(row=r, column=2, value=k + 1)
        ws.cell(row=r, column=3, value=item.get("agenda", ""))
        ws.cell(row=r, column=5, value=item.get("discussion", ""))
        cursor += AGENDA_BLOCK_ROWS

    # 여백 행
    apply_block_style(ws, cursor, spacer_style)
    cursor += 1

    # 지시사항 헤더 배너
    apply_block_style(ws, cursor, instr_header_style)
    ws.merge_cells(start_row=cursor, start_column=2, end_row=cursor, end_column=9)
    ws.row_dimensions[cursor].height = 21.75
    ws.cell(row=cursor, column=2, value=instr_header_values[2])
    cursor += 1

    # 지시사항 표 헤더 (순번/지시사항/담당자/비고)
    apply_block_style(ws, cursor, instr_table_header_style)
    ws.merge_cells(start_row=cursor, start_column=3, end_row=cursor, end_column=7)
    ws.row_dimensions[cursor].height = 18
    for c in cols:
        if instr_table_header_values[c] is not None:
            ws.cell(row=cursor, column=c, value=instr_table_header_values[c])
    cursor += 1

    # 지시사항 데이터 행
    for k, item in enumerate(instructions):
        r = cursor
        apply_block_style(ws, r, instr_row_style)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.row_dimensions[r].height = INSTR_ROW_HEIGHT
        ws.cell(row=r, column=2, value=k + 1)
        ws.cell(row=r, column=3, value=item.get("instruction", ""))
        ws.cell(row=r, column=8, value=item.get("owner", ""))
        ws.cell(row=r, column=9, value=item.get("note", ""))
        cursor += 1

    ws["C5"] = data.get("meeting_name", "")
    ws["G5"] = data.get("datetime", "")
    ws["C6"] = data.get("location", "")
    ws["G6"] = data.get("author", "")
    ws["C7"] = data.get("attendees", "")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)


def send_email_notification(file_path: str, meeting_name: str, to_email: str):
    smtp_user = os.environ["SMTP_USER"]
    smtp_password = os.environ["SMTP_PASSWORD"]

    msg = EmailMessage()
    msg["Subject"] = f"[RP] {meeting_name} 회의록"
    msg["From"] = smtp_user
    msg["To"] = to_email
    msg.set_content(f"'{meeting_name}' 회의록(RP)이 자동 생성되어 첨부파일로 보내드립니다.")

    with open(file_path, "rb") as f:
        msg.add_attachment(
            f.read(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=os.path.basename(file_path),
        )

    with smtplib.SMTP(os.environ.get("SMTP_HOST", "smtp.gmail.com"), int(os.environ.get("SMTP_PORT", "587"))) as server:
        server.starttls()
        server.login(smtp_user, smtp_password)
        server.send_message(msg)


def upload_to_supabase(file_path: str, meeting_name: str, meeting_date: str) -> str:
    supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])

    storage_path = f"{meeting_date}_{safe_filename(meeting_name)}.xlsx"

    with open(file_path, "rb") as f:
        supabase.storage.from_(STORAGE_BUCKET).upload(
            storage_path,
            f,
            {
                "content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "upsert": "true",
            },
        )

    return supabase.storage.from_(STORAGE_BUCKET).get_public_url(storage_path)


def save_to_supabase(data: dict, meeting_date: str, file_url: str):
    supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])
    supabase.table("rp_reports").insert(
        {
            "meeting_name": data.get("meeting_name", ""),
            "meeting_datetime": data.get("datetime", ""),
            "meeting_date": meeting_date,
            "location": data.get("location", ""),
            "author": data.get("author", ""),
            "attendees": data.get("attendees", ""),
            "agenda_items": data.get("agenda_items", []),
            "instructions": data.get("instructions", []),
            "file_url": file_url,
            "source": "rp-automation",
        }
    ).execute()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="전사 텍스트 파일 경로")
    parser.add_argument("--title", default="", help="회의명 (미입력 시 Claude 추출값 사용)")
    parser.add_argument("--date", default=str(date.today()))
    parser.add_argument("--no-upload", action="store_true", help="Supabase 업로드 생략, 로컬 파일만 생성")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        transcript = f.read()

    if os.environ.get("ANTHROPIC_API_KEY"):
        data = call_claude_for_structure(transcript)
    else:
        print("⚠️ ANTHROPIC_API_KEY가 없어 AI 정리 없이 전사록 원문을 그대로 저장합니다.")
        data = build_raw_data(transcript, args.date)

    if args.title:
        data["meeting_name"] = args.title

    output_path = os.path.join(OUTPUT_DIR, f"{args.date}_{safe_filename(data.get('meeting_name', 'RP'))}.xlsx")
    fill_template(data, output_path)
    print(f"✅ 회의록 생성 완료: {output_path}")

    to_email = os.environ.get("RP_EMAIL_TO")
    if to_email and os.environ.get("SMTP_USER") and os.environ.get("SMTP_PASSWORD"):
        send_email_notification(output_path, data.get("meeting_name", "RP"), to_email)
        print(f"✅ 이메일 발송 완료: {to_email}")

    if not args.no_upload and os.environ.get("SUPABASE_URL") and os.environ.get("SUPABASE_SERVICE_KEY"):
        file_url = upload_to_supabase(output_path, data.get("meeting_name", "회의록"), args.date)
        save_to_supabase(data, args.date, file_url)
        print(f"✅ Supabase 업로드 완료: {file_url}")
    elif not args.no_upload:
        print("⚠️ SUPABASE_URL / SUPABASE_SERVICE_KEY가 설정되지 않아 Supabase 업로드를 건너뜁니다.")


if __name__ == "__main__":
    main()
